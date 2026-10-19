from docx import Document


def read_all_text(path):
    doc = Document(path)
    chunks = []

    for p in doc.paragraphs:
        chunks.append(p.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    chunks.append(p.text)

    for section in doc.sections:
        for container in (section.header, section.footer):
            if container is None:
                continue
            for p in container.paragraphs:
                chunks.append(p.text)

    return chunks