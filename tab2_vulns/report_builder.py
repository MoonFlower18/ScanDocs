import os
from docx import Document
from docx.oxml.ns import qn


PH_BDU = "{{BDU}}"
PH_DESC = "{{DESCRIPTION}}"
PH_DATESTAMP = "{{DATESTAMP}}"

ALL_PLACEHOLDERS = [PH_BDU, PH_DESC, PH_DATESTAMP]

def _remove_element(element):
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _delete_paragraph(paragraph):
    _remove_element(paragraph._element)


def _delete_table_row(row):
    _remove_element(row._element)


def _delete_empty_table(table):
    if len(table.rows) == 0:
        _remove_element(table._element)


def _paragraph_text(p):
    return "".join(run.text for run in p.runs)


def _row_text(row):
    parts = []
    for cell in row.cells:
        for p in cell.paragraphs:
            parts.append(_paragraph_text(p))
    return " ".join(parts)


def _row_has_placeholder(row):
    text = _row_text(row)
    return PH_BDU in text or PH_DESC in text


def _replace_in_paragraph(paragraph, placeholder: str, value: str) -> bool:
    p_element = paragraph._element
    t_elements = p_element.findall(".//" + qn("w:t"))
    if not t_elements:
        return False

    texts = [t.text or "" for t in t_elements]
    full_text = "".join(texts)

    if placeholder not in full_text:
        return False

    pos = full_text.find(placeholder)
    end = pos + len(placeholder)

    offset = 0
    first_idx = None
    last_idx = None
    first_offset = 0
    last_offset = 0

    for i, txt in enumerate(texts):
        start = offset
        stop = offset + len(txt)
        if first_idx is None and stop > pos:
            first_idx = i
            first_offset = pos - start
        if last_idx is None and stop >= end:
            last_idx = i
            last_offset = end - start
            break
        offset = stop

    if first_idx is None or last_idx is None:
        return False

    if first_idx == last_idx:
        t_elements[first_idx].text = (
                texts[first_idx][:first_offset]
                + value
                + texts[first_idx][last_offset:]
        )
    else:
        before = texts[first_idx][:first_offset]
        after = texts[last_idx][last_offset:]
        t_elements[first_idx].text = before + value + after
        for i in range(first_idx + 1, last_idx + 1):
            t_elements[i].text = ""

    return True


def _replace_in_paragraphs(paragraphs, placeholder: str, value: str):
    for p in paragraphs:
        _replace_in_paragraph(p, placeholder, value)


def _all_header_footer_containers(doc):
    containers = []
    for section in doc.sections:
        for attr in (
                "header", "footer",
                "first_page_header", "first_page_footer",
                "even_page_header", "even_page_footer",
        ):
            container = getattr(section, attr, None)
            if container is not None:
                containers.append((attr, container))
    return containers


def _scan_tables(tables, placeholder: str, value: str):
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                _replace_in_paragraphs(cell.paragraphs, placeholder, value)
                _scan_tables(cell.tables, placeholder, value)


def _replace_placeholder_everywhere(doc, placeholder: str, value: str):
    if not value:
        return

    _replace_in_paragraphs(doc.paragraphs, placeholder, value)
    _scan_tables(doc.tables, placeholder, value)

    for _name, container in _all_header_footer_containers(doc):
        _replace_in_paragraphs(container.paragraphs, placeholder, value)
        _scan_tables(container.tables, placeholder, value)


def _replace_placeholder_in_row(row, vuln, used):
    replaced = False
    for cell in row.cells:
        for p in cell.paragraphs:
            text = _paragraph_text(p)
            if PH_BDU in text and "bdu" not in used:
                if _replace_in_paragraph(p, PH_BDU, vuln["bdu"]):
                    used.add("bdu")
                    replaced = True
            if PH_DESC in text and "desc" not in used:
                if _replace_in_paragraph(p, PH_DESC, vuln["description"]):
                    used.add("desc")
                    replaced = True
    return replaced


def _delete_row_block(table, start_index):
    rows = list(table.rows)
    i = start_index
    while i < len(rows):
        row = rows[i]
        if i > start_index and _row_has_placeholder(row):
            break
        _delete_table_row(row)
        i += 1


def _cleanup_table(table, filled_rows):
    rows = list(table.rows)
    i = 0
    while i < len(rows):
        row = rows[i]
        if _row_has_placeholder(row) and i not in filled_rows:
            _delete_row_block(table, i)
            rows = list(table.rows)
            continue
        i += 1


def find_unreplaced(doc) -> list:
    leftovers = set()

    def scan(paragraphs):
        for p in paragraphs:
            text = "".join(r.text for r in p.runs)
            for ph in ALL_PLACEHOLDERS:
                if ph in text:
                    leftovers.add(ph)

    def scan_tables(tables):
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    scan(cell.paragraphs)
                    scan_tables(cell.tables)

    scan(doc.paragraphs)
    scan_tables(doc.tables)

    for _name, container in _all_header_footer_containers(doc):
        scan(container.paragraphs)
        scan_tables(container.tables)

    return sorted(leftovers)


def build_report(
        template_path: str,
        output_path: str,
        vulnerabilities: list,
        datestamp: str = "",
) -> list:

    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Шаблон не найден: {template_path}")

    doc = Document(template_path)

    _replace_placeholder_everywhere(doc, PH_DATESTAMP, datestamp)

    filled_rows = set()
    vuln_idx = 0

    for table in doc.tables:
        rows = list(table.rows)
        for idx, row in enumerate(rows):
            if vuln_idx >= len(vulnerabilities):
                break
            if not _row_has_placeholder(row):
                continue
            used = set()
            if _replace_placeholder_in_row(row, vulnerabilities[vuln_idx], used):
                filled_rows.add(idx)
                vuln_idx += 1

    for table in doc.tables:
        _cleanup_table(table, filled_rows)
        _delete_empty_table(table)

    for vuln in vulnerabilities[vuln_idx:]:
        p = doc.add_paragraph()
        p.add_run(f"{vuln['bdu']} — {vuln['description']}")

    leftovers = find_unreplaced(doc)

    doc.save(output_path)
    return leftovers